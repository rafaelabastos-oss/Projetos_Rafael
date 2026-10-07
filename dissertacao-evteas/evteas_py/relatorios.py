"""EVTEAS-Py — módulo relatorios.

Resumo executivo, tabelas no formato da dissertação, marcadores de resultados ([R-xx]) para
inserção no texto, figuras, workbook Excel, JSON, relatório Markdown e preenchimento automático
de um arquivo .docx que contenha os marcadores (Etapa 12 da DSR — comunicação dos resultados).
"""
from __future__ import annotations

import json
import math
import os
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd

from .config import EVTEASConfig, VERSAO, fmt_num, fmt_pct, fmt_rs
from .entradas import marcadores_entradas, tabela_entradas
from .incerteza import rotulo
from .ponderacao import KPIS, NOMES_DIMENSOES

DIAG_TIR = {"convencional": "convencional", "multiplas": "múltiplas", "multiplas_possiveis": "múltiplas possíveis",
            "sem_solucao": "sem solução"}


def _f(x) -> float:
    try:
        v = np.atleast_1d(np.asarray(x, dtype=float))[0]
        return float(v)
    except (TypeError, ValueError):
        return float("nan")


def _mes(v) -> str:
    return "não recupera no horizonte" if not math.isfinite(_f(v)) else fmt_num(_f(v), 1)


# ---------------------------------------------------------------------------------------------
# Marcadores de resultados para o texto da dissertação
# ---------------------------------------------------------------------------------------------
def marcadores(resultado: Dict[str, Any], cfg: EVTEASConfig, alternativas: Optional[pd.DataFrame] = None,
               testes: Optional[pd.DataFrame] = None,
               configs_alternativas: Optional[Dict[str, EVTEASConfig]] = None) -> pd.DataFrame:
    """Tabela Marcador | Descrição | Valor | Seção.

    No texto da dissertação, [E-xx] são as entradas informadas pelo usuário e [R-xx] os resultados
    da simulação; ambos são preenchidos depois da execução, com os valores desta tabela.
    """
    T, E, A, S = resultado["tecnico"], resultado["economico"], resultado["ambiental"], resultado["social"]
    G, LG, I, D = resultado["governanca"], resultado["lean_green"], resultado["indice"], resultado["decisao"]
    e, t = cfg.economico, cfg.tecnico
    L: List[tuple] = []
    add = lambda m, desc, val, sec: L.append((m, desc, val, sec))

    # Entradas informadas pelo usuário (Quadro de premissas, Seção 4.2) — marcadores [E-xx]
    for linha in marcadores_entradas(cfg, configs_alternativas):
        L.append(linha)

    # Técnica (4.3)
    add("R-T01", "Restrição ativa da biomassa", T["restricao_ativa"], "4.3")
    add("R-T02", "Juvenis estocados por ciclo", fmt_num(_f(T["alevinos_por_ciclo"]), 0), "4.3")
    add("R-T03", "Densidade de estocagem (peixes/m³)", fmt_num(_f(T["densidade_peixes_m3"]), 2), "4.3")
    add("R-T04", "Despesca por ciclo (kg)", fmt_num(_f(T["biomassa_despesca_ciclo_kg"]), 0), "4.3")
    add("R-T05", "Produção anual (kg/ano)", fmt_num(_f(T["producao_ano_kg"]), 0), "4.3")
    add("R-T06", "Produtividade por área (kg/m²/ano)", fmt_num(_f(T["produtividade_kg_m2_ano"]), 2), "4.3")
    add("R-T07", "Produtividade relativa ao potencial do sistema", fmt_pct(_f(T["produtividade_relativa"])), "4.3")
    add("R-T08", "OEE aquícola", fmt_num(_f(T["oee"]), 3), "4.3")
    add("R-T09", "Disponibilidade; performance; sobrevivência",
        f"{fmt_num(_f(T['disponibilidade']), 4)}; {fmt_num(_f(T['performance']), 2)}; {fmt_num(_f(T['sobrevivencia']), 2)}", "4.3")
    add("R-T10", "Consumo de ração (kg/ano)", fmt_num(_f(T["racao_ano_kg"]), 0), "4.3")
    add("R-T11", "Desperdício: ração excedente (R$/ano)", fmt_rs(_f(LG["racao_excedente_rs"])), "4.3")
    add("R-T12", "Desperdício: mortalidade evitável (R$/ano)", fmt_rs(_f(LG["mortalidade_evitavel_rs"])), "4.3")
    add("R-T13", "Desperdício: energia excedente (R$/ano)", fmt_rs(_f(LG["energia_excedente_rs"])), "4.3")
    add("R-T14", "Desperdícios Lean-Green totais (R$/ano)", fmt_rs(_f(LG["desperdicio_total_rs"])), "4.3")
    add("R-T15", "Desperdícios em % do custo operacional", fmt_pct(_f(LG["desperdicio_pct_custo_operacional"])), "4.3")

    # Econômica (4.4)
    add("R-E01", "Meses até a primeira despesca", str(E["meses_ate_despesca"]), "4.4")
    add("R-E02", "Investimento inicial do projeto (CAPEX + capital de giro)", fmt_rs(_f(E["investimento"])), "4.4")
    add("R-E03", "Investimento do beneficiário (com fomento)", fmt_rs(_f(E["investimento_beneficiario"])), "4.4")
    add("R-E04", "VPL do projeto", fmt_rs(_f(E["vpl"])), "4.4")
    add("R-E05", "VPL do beneficiário", fmt_rs(_f(E["vpl_beneficiario"])), "4.4")
    add("R-E06", "TIR anual do projeto (diagnóstico)", f"{fmt_pct(_f(E['tir_aa']))} ({DIAG_TIR.get(str(E['diagnostico_tir']), E['diagnostico_tir'])})", "4.4")
    add("R-E07", "TIR anual do beneficiário", fmt_pct(_f(E["tir_aa_beneficiario"])), "4.4")
    add("R-E08", "Payback descontado do projeto (meses)", _mes(E["payback_descontado_meses"]), "4.4")
    add("R-E09", "Payback descontado do beneficiário (meses)", _mes(E["payback_descontado_beneficiario_meses"]), "4.4")
    add("R-E10", "Payback simples do projeto (meses)", _mes(E["payback_simples_meses"]), "4.4")
    add("R-E11", "Índice de lucratividade", fmt_num(_f(E["indice_lucratividade"]), 3), "4.4")
    add("R-E12", f"Receita bruta no ano de regime (ano {e.ano_regime})", fmt_rs(_f(E["receita_regime"])), "4.4")
    add("R-E13", "Resultado (sobras) líquido no ano de regime", fmt_rs(_f(E["lucro_liquido_regime"])), "4.4")
    add("R-E14", "ROI anual em regime", fmt_pct(_f(E["roi_regime"])), "4.4")
    add("R-E15", "Ponto de equilíbrio (kg/ano)", fmt_num(_f(E["ponto_equilibrio_kg_ano"]), 0), "4.4")
    add("R-E16", "Margem de segurança", fmt_pct(_f(E["margem_seguranca"])), "4.4")
    add("R-E17", "Custo total de produção (R$/kg)", fmt_num(_f(E["custo_total_kg"]), 2), "4.4")
    add("R-E18", "Custo operacional efetivo (R$/kg)", fmt_num(_f(E["custo_operacional_kg"]), 2), "4.4")
    add("R-E19", "Participação da ração no custo operacional", fmt_pct(_f(E["participacao_racao_custo_op"])), "4.4")
    add("R-E20", "TMA (% a.a.)", fmt_num(e.tma_aa_pct, 1) + "%", "4.4")

    # Ambiental (4.5)
    add("R-A01", "Água captada (m³/ano)", fmt_num(_f(A["agua_captada_m3_ano"]), 0), "4.5")
    add("R-A02", "Índice de uso da água (kg/m³)", fmt_num(_f(A["indice_uso_agua_kg_m3"]), 3), "4.5")
    add("R-A03", "Pegada hídrica azul (m³/t)", fmt_num(_f(A["ph_azul_m3_t"]), 0), "4.5")
    add("R-A04", "Pegada hídrica cinza (m³/t)", fmt_num(_f(A["ph_cinza_m3_t"]), 0), "4.5")
    add("R-A05", "Poluente crítico da pegada cinza", str(np.atleast_1d(A["poluente_critico"])[0]), "4.5")
    add("R-A06", "Fósforo aportado pela ração (kg/ano)", fmt_num(_f(A["p_aportado_kg"]), 0), "4.5")
    add("R-A07", "Fósforo retido na biomassa (kg/ano)", fmt_num(_f(A["p_retido_kg"]), 0), "4.5")
    add("R-A08", "Fósforo lançado após tratamento (kg/ano)", fmt_num(_f(A["p_lancado_kg"]), 0), "4.5")
    add("R-A09", "Nitrogênio lançado após tratamento (kg/ano)", fmt_num(_f(A["n_lancado_kg"]), 0), "4.5")
    add("R-A10", "Emissões totais (tCO2e/ano)", fmt_num(_f(A["co2_total_kg"]) / 1000, 1), "4.5")
    add("R-A11", "Emissões: ração / energia / diesel (tCO2e)",
        f"{fmt_num(_f(A['co2_racao_kg']) / 1000, 1)} / {fmt_num(_f(A['co2_energia_kg']) / 1000, 1)} / {fmt_num(_f(A['co2_diesel_kg']) / 1000, 1)}", "4.5")
    add("R-A12", "Participação da ração nas emissões", fmt_pct(_f(A["co2_racao_kg"]) / _f(A["co2_total_kg"])), "4.5")
    add("R-A13", "Intensidade de carbono (kgCO2e/kg)", fmt_num(_f(A["intensidade_carbono_kg_kg"]), 2), "4.5")
    add("R-A14", "Ecoeficiência (R$ de valor adicionado/kgCO2e)", fmt_num(_f(A["ecoeficiencia_rs_kgco2e"]), 2), "4.5")
    add("R-A15", "Score de conformidade ambiental (0–10)", fmt_num(_f(A["score_conformidade_0_10"]), 2), "4.5")
    nao = A["checklist"][A["checklist"]["Atendido"] == "não"]["Item"].tolist()
    add("R-A16", "Itens do checklist ambiental não atendidos", "; ".join(nao) if nao else "nenhum", "4.5")
    add("R-A17", "Itens eliminatórios violados", "; ".join(A["vetos_ambientais"]) if A["vetos_ambientais"] else "nenhum", "4.5")

    # Social e ESG (4.6)
    add("R-S01", "Retenção de Valor Local (RVL)", fmt_pct(_f(S["rvl_pct"]) / 100), "4.6")
    add("R-S02", "Índice de Impacto Local (renda local/receita)", fmt_pct(_f(S["indice_impacto_local"])), "4.6")
    add("R-S03", "Renda mensal média por trabalhador", fmt_rs(_f(S["renda_mensal_trabalhador"])), "4.6")
    add("R-S04", "Renda em salários mínimos", fmt_num(_f(S["renda_salarios_minimos"]), 2), "4.6")
    add("R-S05", "Licença Social para Operar (0–100)", fmt_num(S["lso_0_100"], 1), "4.6")
    add("R-S06", "Score de risco social (0–1)", fmt_num(S["risco_social"], 3), "4.6")
    add("R-S07", "Radar Social (média dos seis eixos)", fmt_num(_f(S["radar_social_media"]), 3), "4.6")
    radar = {k: _f(v) for k, v in S["radar"].items()}
    menores = sorted(radar, key=radar.get)[:2]
    add("R-S08", "Eixos de menor valor no Radar Social", " e ".join(m.lower() for m in menores), "4.6")
    ods = resultado["ods"]
    atend = ods[ods["Atendido"] == "sim"]["ODS"].tolist()
    nao_ods = ods[ods["Atendido"] == "não"]["ODS"].tolist()
    add("R-S09", "ODS atendidos", f"{len(atend)} de {len(ods)} ({', '.join(atend)})", "4.6")
    add("R-S10", "ODS não atendidos", ", ".join(nao_ods) if nao_ods else "nenhum", "4.6")
    add("R-S11", "Governança: itens atendidos e score", f"{G['itens_atendidos']} de {G['itens_total']} ({fmt_num(G['score_governanca'], 3)})", "4.6")
    pil = I["pilares_esg"]
    add("R-S12", "Pilares ESG: ambiental / social / governança",
        f"{fmt_num(_f(pil['ambiental']), 3)} / {fmt_num(_f(pil['social']), 3)} / {fmt_num(_f(pil['governanca']), 3)}", "4.6")
    add("R-S13", "Score ESG", fmt_num(_f(I["score_esg"]), 3), "4.6")

    # Índice e decisão (4.7)
    sc = I["scores"]
    pw = I["pesos"]["dimensoes"]
    add("R-I01", "Método e pesos dimensionais (T/E/A/S)",
        f"{I['pesos']['metodo']}: {fmt_num(pw['tecnica'], 3)} / {fmt_num(pw['economica'], 3)} / {fmt_num(pw['ambiental'], 3)} / {fmt_num(pw['social'], 3)}", "4.7")
    add("R-I02", "Score técnico", fmt_num(_f(sc["tecnica"]), 3), "4.7")
    add("R-I03", "Score econômico", fmt_num(_f(sc["economica"]), 3), "4.7")
    add("R-I04", "Score ambiental", fmt_num(_f(sc["ambiental"]), 3), "4.7")
    add("R-I05", "Score social", fmt_num(_f(sc["social"]), 3), "4.7")
    add("R-I06", "Índice EVTEAS", fmt_num(_f(I["indice_evteas"]), 3), "4.7")
    norm = resultado["normalizacao"].sort_values("Normalizado")
    add("R-I07", "KPIs de pior desempenho relativo", " e ".join(norm["KPI"].head(2).tolist()), "4.7")
    add("R-I08", "Classificação final", D["classificacao"], "4.7")
    add("R-I09", "Vetos aplicados", "; ".join(D["vetos"]) if D["vetos"] else "nenhum", "4.7")
    add("R-I10", "Ressalvas", "; ".join(D["ressalvas"]) if D["ressalvas"] else "nenhuma", "4.7")
    if I["pesos"]["metodo"] == "likert":
        add("R-I11", "Concordância média entre avaliadores (Likert)", fmt_num(I["pesos"]["concordancia_media"], 3), "4.7")
    if I["pesos"]["metodo"] == "ahp":
        add("R-I11", "Razão de consistência AHP", fmt_num(I["pesos"]["razao_consistencia"], 3), "4.7")

    # Monte Carlo (4.8)
    if "monte_carlo" in resultado:
        mc = resultado["monte_carlo"]
        st = mc["estatisticas"]
        add("R-M01", "Iterações e semente", f"{fmt_num(mc['n'], 0)} (seed {mc['seed']})", "4.8")
        add("R-M02", "Média do VPL (IC 95% da média)", f"{fmt_rs(st['media'])} ({fmt_rs(st['ic_inf'])} a {fmt_rs(st['ic_sup'])})", "4.8")
        add("R-M03", "Mediana; desvio-padrão; coeficiente de variação",
            f"{fmt_rs(st['mediana'])}; {fmt_rs(st['desvio_padrao'])}; {fmt_num(st['coef_variacao'], 2)}", "4.8")
        add("R-M04", "Percentis P5 / P50 / P95 do VPL", f"{fmt_rs(st['p5'])} / {fmt_rs(st['p50'])} / {fmt_rs(st['p95'])}", "4.8")
        add("R-M05", "VaR 5% e CVaR 5% do VPL", f"{fmt_rs(st['var_5'])} e {fmt_rs(st['cvar_5'])}", "4.8")
        add("R-M06", "P(VPL > 0) = P(TIR > TMA)", fmt_pct(st["prob_vpl_positivo"]), "4.8")
        add("R-M07", "TIR anual: P5 / P50 / P95", f"{fmt_pct(st['tir_p5'])} / {fmt_pct(st['tir_p50'])} / {fmt_pct(st['tir_p95'])}", "4.8")
        add("R-M08", "Índice EVTEAS: média; P(índice ≥ limiar)", f"{fmt_num(st['indice_media'], 3)}; {fmt_pct(st['prob_indice_viavel'])}", "4.8")
        add("R-M09", "VPL do beneficiário: média; P(VPL > 0)", f"{fmt_rs(st['vpl_beneficiario_media'])}; {fmt_pct(st['prob_vpl_beneficiario_positivo'])}", "4.8")
        add("R-M10", "Erro-padrão da média do VPL", f"{fmt_rs(st['erro_padrao'])} ({fmt_pct(st['erro_padrao'] / st['desvio_padrao'] if st['desvio_padrao'] else float('nan'))} do desvio-padrão)", "4.8")
        imp = mc["importancia"]
        add("R-M11", "Três variáveis mais importantes (Spearman)",
            "; ".join(f"{r['Rótulo']} ({fmt_num(r['Spearman (VPL)'], 3)})" for _, r in imp.head(3).iterrows()), "4.9")

    # Sensibilidade, valores críticos e cenários (4.9)
    if "sensibilidade" in resultado:
        sens = resultado["sensibilidade"]
        top = sens.iloc[0]
        add("R-V01", "Variável de maior impacto (±10%)", top["Rótulo"], "4.9")
        add("R-V02", "Variação do VPL para ±10% na variável de maior impacto", fmt_rs(abs(top["ΔVPL (+10%)"])), "4.9")
        vc = resultado["valores_criticos"].set_index("Variável")
        if "economico.preco_venda_kg" in vc.index:
            p = vc.loc["economico.preco_venda_kg"]
            add("R-V03", "Preço crítico (VPL = 0) e folga", f"{fmt_rs(p['Valor crítico (VPL = 0)'], 2)}/kg ({fmt_pct(p['Folga relativa'])})", "4.9")
        if "tecnico.fcr" in vc.index:
            p = vc.loc["tecnico.fcr"]
            add("R-V04", "FCR crítico e folga", f"{fmt_num(p['Valor crítico (VPL = 0)'], 3)} ({fmt_pct(p['Folga relativa'])})", "4.9")
        if "tecnico.mortalidade_pct" in vc.index:
            p = vc.loc["tecnico.mortalidade_pct"]
            add("R-V05", "Mortalidade crítica", f"{fmt_num(p['Valor crítico (VPL = 0)'], 1)}%", "4.9")
        cen = resultado["cenarios"].set_index("Cenário")
        add("R-V06", "Cenário pessimista: VPL; índice", f"{fmt_rs(cen.loc['Pessimista', 'VPL'])}; {fmt_num(cen.loc['Pessimista', 'Índice EVTEAS'], 3)}", "4.9")
        add("R-V07", "Cenário otimista: VPL; índice", f"{fmt_rs(cen.loc['Otimista', 'VPL'])}; {fmt_num(cen.loc['Otimista', 'Índice EVTEAS'], 3)}", "4.9")

    # Alternativas (4.9)
    if alternativas is not None and len(alternativas):
        alt = alternativas
        best = alt.iloc[0]
        add("R-X01", "Ordem TOPSIS das alternativas", " > ".join(alt["Alternativa"].tolist()), "4.9")
        add("R-X02", "Alternativa preferida e proximidade à solução ideal", f"{best['Alternativa']} ({fmt_num(best['TOPSIS'], 3)})", "4.9")
        ix = alt.set_index("Alternativa")
        for nome in ix.index:
            add(f"R-X{nome}", f"Alternativa {nome}: VPL; P(VPL>0); PH cinza; kgCO2e/kg",
                f"{fmt_rs(ix.loc[nome, 'VPL'])}; {fmt_pct(ix.loc[nome, 'P(VPL>0)'])}; {fmt_num(ix.loc[nome, 'PH cinza (m³/t)'], 0)} m³/t; {fmt_num(ix.loc[nome, 'kgCO2e/kg'], 2)}", "4.9")
        add("R-X03", "Maior probabilidade de VPL positivo entre as alternativas",
            f"{fmt_pct(alt['P(VPL>0)'].max())} (alternativa {alt.loc[alt['P(VPL>0)'].idxmax(), 'Alternativa']})", "4.9")

    # Verificação (4.10) e execução (4.11)
    vv = resultado["vv"]
    add("R-W01", "Invariantes aprovados", f"{vv['aprovadas']} de {vv['total']}", "4.10")
    if testes is not None and len(testes):
        add("R-W02", "Testes automatizados aprovados", f"{int(testes['Aprovado'].sum())} de {len(testes)}", "4.10")
    add("R-W03", "Tempo de execução do estudo (s)", fmt_num(resultado.get("tempo_execucao_s", float("nan")), 1), "4.11")
    add("R-W04", "Versão do EVTEAS-Py", VERSAO, "4.11")

    for chave, texto in textos_interpretativos(resultado, cfg, alternativas).items():
        add(chave, "Texto interpretativo sugerido (revise antes de usar)", texto, "4.x")
    return pd.DataFrame(L, columns=["Marcador", "Descrição", "Valor", "Seção"])


def textos_interpretativos(resultado: Dict[str, Any], cfg: EVTEASConfig,
                           alternativas: Optional[pd.DataFrame] = None) -> Dict[str, str]:
    """Frases geradas por regras explícitas, para apoiar (não substituir) a interpretação."""
    E, D = resultado["economico"], resultado["decisao"]
    tma = cfg.economico.tma_aa_pct / 100
    vpl, tir_aa = _f(E["vpl"]), _f(E["tir_aa"])
    txt = {}
    if vpl >= 0:
        txt["R-TXT01"] = (f"Na perspectiva do projeto, o VPL de {fmt_rs(vpl)} e a TIR de {fmt_pct(tir_aa)} a.a., superior "
                          f"à TMA de {fmt_pct(tma)} a.a., indicam viabilidade econômica no cenário determinístico.")
    else:
        txt["R-TXT01"] = (f"Na perspectiva do projeto, o VPL de {fmt_rs(vpl)} e a TIR de {fmt_pct(tir_aa)} a.a., inferior "
                          f"à TMA de {fmt_pct(tma)} a.a., indicam que o projeto não remunera o capital à taxa mínima de atratividade no cenário determinístico.")
    if "monte_carlo" in resultado:
        st = resultado["monte_carlo"]["estatisticas"]
        rel = "abaixo" if st["media"] < vpl else "acima"
        txt["R-TXT02"] = (f"A média do VPL simulado ({fmt_rs(st['media'])}) ficou {rel} do VPL calculado na moda das premissas "
                          f"({fmt_rs(vpl)}), e a probabilidade de VPL positivo foi de {fmt_pct(st['prob_vpl_positivo'])}"
                          + (", o que evidencia o efeito das distribuições assimétricas desfavoráveis." if rel == "abaixo"
                             else ", o que indica distribuições com caudas favoráveis."))
        txt["R-TXT03"] = (f"Na perspectiva do beneficiário, com fomento de {fmt_num(cfg.economico.fomento_nao_reembolsavel_pct, 0)}% do CAPEX, "
                          f"a probabilidade de VPL positivo passa a {fmt_pct(st['prob_vpl_beneficiario_positivo'])}.")
    if "sensibilidade" in resultado:
        s = resultado["sensibilidade"]
        rot = [r[0].lower() + r[1:] for r in s["Rótulo"].head(3)]
        txt["R-TXT04"] = ("As variáveis de maior impacto sobre o VPL foram " + ", ".join(rot[:-1]) + " e " + rot[-1]
                          + ", que concentram as prioridades de coleta de dados e de gestão de risco.")
    txt["R-TXT05"] = (f"O estudo foi classificado como \"{D['classificacao']}\""
                      + (f", em razão de: {'; '.join(D['vetos'])}" if D["vetos"] else "")
                      + (f"; ressalvas: {'; '.join(D['ressalvas'])}" if D["ressalvas"] else "") + ".")
    if alternativas is not None and len(alternativas):
        b = alternativas.iloc[0]
        txt["R-TXT06"] = (f"A alternativa {b['Alternativa']} obteve a maior proximidade à solução ideal ({fmt_num(b['TOPSIS'], 3)}), "
                          f"com VPL de {fmt_rs(b['VPL'])} e probabilidade de VPL positivo de {fmt_pct(b['P(VPL>0)'])}.")
    return txt


# ---------------------------------------------------------------------------------------------
# Tabelas no formato da dissertação
# ---------------------------------------------------------------------------------------------
def tabelas_dissertacao(resultado: Dict[str, Any], cfg: EVTEASConfig,
                        alternativas: Optional[pd.DataFrame] = None) -> Dict[str, pd.DataFrame]:
    m = marcadores(resultado, cfg, alternativas).set_index("Marcador")["Valor"]
    v = lambda k: m.get(k, "n/d")
    tab = {}
    tab["Tabela - Indicadores econômicos"] = pd.DataFrame([
        ["Investimento inicial (CAPEX + capital de giro)", v("R-E02"), v("R-E03")],
        [f"VPL (TMA de {v('R-E20')} a.a.)", v("R-E04"), v("R-E05")],
        ["TIR anual (diagnóstico)", v("R-E06"), v("R-E07")],
        ["Payback descontado (meses)", v("R-E08"), v("R-E09")],
        ["Payback simples (meses)", v("R-E10"), "—"],
        ["Índice de lucratividade", v("R-E11"), "—"],
        [f"Receita bruta no ano de regime (ano {cfg.economico.ano_regime})", v("R-E12"), "—"],
        ["Resultado (sobras) líquido no ano de regime", v("R-E13"), "—"],
        ["ROI anual em regime (Quadro 5)", v("R-E14"), "—"],
        ["Ponto de equilíbrio (kg/ano) e margem de segurança", f"{v('R-E15')} kg; {v('R-E16')}", "—"],
        ["Custo total e custo operacional efetivo (R$/kg)", f"{v('R-E17')} e {v('R-E18')}", "—"],
    ], columns=["Indicador", "Projeto", "Beneficiário (com fomento)"])
    tab["Tabela - Indicadores ambientais, sociais e de governança"] = pd.DataFrame([
        ["Ambiental", "Água captada / índice de uso da água", f"{v('R-A01')} m³/ano; {v('R-A02')} kg/m³"],
        ["Ambiental", "Pegada hídrica azul", f"{v('R-A03')} m³/t"],
        ["Ambiental", "Pegada hídrica cinza (poluente crítico)", f"{v('R-A04')} m³/t ({v('R-A05')})"],
        ["Ambiental", "N e P lançados após tratamento", f"{v('R-A09')} kg N; {v('R-A08')} kg P"],
        ["Ambiental", "Emissões (ração / energia / diesel)", f"{v('R-A10')} tCO2e/ano ({v('R-A11')})"],
        ["Ambiental", "Intensidade de carbono; ecoeficiência", f"{v('R-A13')} kgCO2e/kg; R$ {v('R-A14')}/kgCO2e"],
        ["Ambiental", "Score de conformidade ambiental", f"{v('R-A15')} de 10"],
        ["Social", "Retenção de Valor Local (RVL)", f"{v('R-S01')} do custo operacional"],
        ["Social", "Índice de Impacto Local", f"{v('R-S02')} da receita retorna como renda local"],
        ["Social", "Renda mensal por trabalhador", f"{v('R-S03')} ({v('R-S04')} salários mínimos)"],
        ["Social", "LSO; score de risco social", f"{v('R-S05')} de 100; {v('R-S06')}"],
        ["Social", "Radar Social (média dos seis eixos)", v("R-S07")],
        ["Governança", "Checklist do pilar G", v("R-S11")],
    ], columns=["Dimensão", "Indicador", "Resultado"])
    norm = resultado["normalizacao"].copy()
    norm["Valor bruto"] = norm["Valor bruto"].map(lambda x: fmt_num(x, 3))
    norm["Normalizado"] = norm["Normalizado"].map(lambda x: fmt_num(x, 3))
    tab["Quadro - Normalização dos KPIs"] = norm
    if "monte_carlo" in resultado:
        tab["Tabela - Estatísticas da simulação de Monte Carlo"] = pd.DataFrame([
            ["Média do VPL (IC 95% da média)", v("R-M02")], ["Mediana; desvio-padrão; coeficiente de variação", v("R-M03")],
            ["Percentis P5 / P50 / P95", v("R-M04")], ["VaR 5% e CVaR 5% do VPL", v("R-M05")],
            ["P(VPL > 0) = P(TIR > TMA)", v("R-M06")], ["TIR anual: P5 / P50 / P95", v("R-M07")],
            ["Índice EVTEAS: média; P(índice ≥ 0,60)", v("R-M08")], ["VPL do beneficiário: média; P(VPL > 0)", v("R-M09")],
            ["Erro-padrão da média", v("R-M10")],
        ], columns=["Estatística", "Valor"])
    if "sensibilidade" in resultado:
        s = resultado["sensibilidade"].copy()
        vc = resultado["valores_criticos"].set_index("Variável")
        imp = resultado.get("monte_carlo", {}).get("importancia")
        rho = imp.set_index("Variável")["Spearman (VPL)"] if imp is not None else pd.Series(dtype=float)
        linhas = []
        for _, r in s.iterrows():
            var = r["Variável"]
            crit = "—"
            if var in vc.index and np.isfinite(vc.loc[var, "Valor crítico (VPL = 0)"]):
                crit = f"{fmt_num(vc.loc[var, 'Valor crítico (VPL = 0)'], 3)} ({fmt_pct(vc.loc[var, 'Folga relativa'])})"
            linhas.append([r["Rótulo"], fmt_rs(r["ΔVPL (-10%)"]), fmt_rs(r["ΔVPL (+10%)"]), crit,
                           fmt_num(rho[var], 3) if var in rho.index else "—"])
        tab["Tabela - Sensibilidade, valores críticos e importância"] = pd.DataFrame(
            linhas, columns=["Variável", "ΔVPL (−10%)", "ΔVPL (+10%)", "Valor crítico (VPL = 0)", "Spearman (MC)"])
    if alternativas is not None and len(alternativas):
        a = alternativas.copy()
        tab["Tabela - Comparação de alternativas por TOPSIS"] = pd.DataFrame({
            "Alternativa": a["Alternativa"], "VPL": a["VPL"].map(fmt_rs), "P(VPL>0)": a["P(VPL>0)"].map(fmt_pct),
            "PH cinza (m³/t)": a["PH cinza (m³/t)"].map(lambda x: fmt_num(x, 0)),
            "kgCO2e/kg": a["kgCO2e/kg"].map(lambda x: fmt_num(x, 2)), "RVL (%)": a["RVL (%)"].map(lambda x: fmt_num(x, 1)),
            "Índice": a["Índice"].map(lambda x: fmt_num(x, 3)), "TOPSIS": a["TOPSIS"].map(lambda x: fmt_num(x, 3)),
            "Posição": a["Posição"]})
    return tab


# ---------------------------------------------------------------------------------------------
# Figuras
# ---------------------------------------------------------------------------------------------
FIGURAS = {
    "FIG-CRESCIMENTO": ("fig_curva_crescimento.png", "Curva de crescimento, biomassa e ração acumulada no ciclo"),
    "FIG-ECONOMICO": ("fig_receita_custos_fluxo.png", "Receita, custos operacionais e fluxo de caixa acumulado"),
    "FIG-RADAR": ("fig_radar_social.png", "Radar Social"),
    "FIG-INDICE": ("fig_scores_indice.png", "Scores dimensionais e índice EVTEAS"),
    "FIG-VPL": ("fig_distribuicao_vpl.png", "Distribuição e função de probabilidade acumulada do VPL"),
    "FIG-CONVERGENCIA": ("fig_convergencia.png", "Convergência da média do VPL com intervalo de confiança de 95%"),
    "FIG-TORNADO": ("fig_tornado.png", "Diagrama de tornado da sensibilidade do VPL a ±10%"),
    "FIG-TOPSIS": ("fig_topsis.png", "Proximidade relativa à solução ideal das alternativas de projeto"),
    "FIG-LEANGREEN": ("fig_lean_green.png", "Desperdícios Lean-Green evitáveis"),
}


def _mpl():
    import matplotlib
    if "inline" not in matplotlib.get_backend().lower():
        try:
            matplotlib.use("Agg")
        except Exception:  # noqa: BLE001
            pass
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 10, "axes.grid": True, "grid.alpha": 0.3, "figure.dpi": 100,
                         "axes.spines.top": False, "axes.spines.right": False})
    return plt


def gerar_figuras(resultado: Dict[str, Any], cfg: EVTEASConfig, pasta: str,
                  alternativas: Optional[pd.DataFrame] = None, mostrar: bool = False) -> Dict[str, str]:
    """Gera as figuras em PNG (300 dpi) e devolve {marcador: caminho}."""
    plt = _mpl()
    from matplotlib.ticker import FuncFormatter
    Path(pasta).mkdir(parents=True, exist_ok=True)
    from .entradas import fmt_auto
    reais = FuncFormatter(lambda x, _: fmt_num(x / 1000, 0) + " mil")
    virg = FuncFormatter(lambda x, _: fmt_auto(round(float(x), 6)))   # vírgula decimal (pt-BR)
    saida = {}

    def salvar(fig, chave):
        caminho = os.path.join(pasta, FIGURAS[chave][0])
        for ax_ in fig.axes:
            if getattr(ax_, "name", "") == "polar":
                continue
            for eixo in (ax_.xaxis, ax_.yaxis):
                if not isinstance(eixo.get_major_formatter(), FuncFormatter) and eixo.get_scale() == "linear":
                    eixo.set_major_formatter(virg)
        fig.tight_layout()
        fig.savefig(caminho, dpi=300, bbox_inches="tight")
        if mostrar:
            plt.show()
        plt.close(fig)
        saida[chave] = caminho

    c = resultado["curva_crescimento"]
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.plot(c["Dia"], c["Peso médio (g)"], color="#1f77b4", label="Peso médio (g)")
    ax.set_xlabel("Dia do ciclo")
    ax.set_ylabel("Peso médio (g)")
    ax2 = ax.twinx()
    ax2.plot(c["Dia"], c["Biomassa (kg)"], color="#2ca02c", label="Biomassa (kg)")
    ax2.plot(c["Dia"], c["Ração acumulada (kg)"], color="#ff7f0e", ls="--", label="Ração acumulada (kg)")
    ax2.set_ylabel("kg")
    ax2.grid(False)
    linhas = ax.get_lines() + ax2.get_lines()
    ax.legend(linhas, [l.get_label() for l in linhas], loc="upper left")
    salvar(fig, "FIG-CRESCIMENTO")

    dre = resultado["dre_anual"]
    fluxo = np.asarray(resultado["fluxo_caixa"])
    acum = np.cumsum(fluxo)
    fig, ax = plt.subplots(figsize=(8, 4))
    anos = dre["Ano"].to_numpy()
    ax.bar(anos - 0.7, dre["Receita Bruta"], width=0.4, label="Receita bruta (ano)", color="#4c9be8")
    ax.bar(anos - 0.3, dre["Custos Variáveis"] + dre["Custos Fixos"], width=0.4, label="Custos operacionais (ano)",
           color="#f2a65a")
    ax.plot(np.arange(len(acum)) / 12, acum, color="#333333", label="Fluxo de caixa acumulado")
    ax.axhline(0, color="black", lw=0.8)
    ax.yaxis.set_major_formatter(reais)
    ax.set_xlabel("Ano")
    ax.set_ylabel("R$")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.15), ncol=3, frameon=False)
    salvar(fig, "FIG-ECONOMICO")

    radar = {k: _f(v) for k, v in resultado["social"]["radar"].items()}
    nomes = list(radar)
    vals = list(radar.values()) + [list(radar.values())[0]]
    ang = np.linspace(0, 2 * np.pi, len(nomes), endpoint=False).tolist()
    ang += ang[:1]
    fig = plt.figure(figsize=(6, 5))
    ax = fig.add_subplot(111, polar=True)
    ax.plot(ang, vals, color="#2ca02c")
    ax.fill(ang, vals, color="#2ca02c", alpha=0.25)
    ax.set_xticks(ang[:-1])
    ax.set_xticklabels(nomes)
    ax.set_ylim(0, 1)
    salvar(fig, "FIG-RADAR")

    sc = resultado["indice"]["scores"]
    fig, ax = plt.subplots(figsize=(8, 4))
    nomes = [NOMES_DIMENSOES[d] for d in sc] + ["Índice EVTEAS"]
    vals = [_f(v) for v in sc.values()] + [_f(resultado["indice"]["indice_evteas"])]
    cores = ["#4c9be8"] * len(sc) + ["#2ca02c"]
    ax.bar(nomes, vals, color=cores)
    ax.axhline(cfg.decisao.limiar_indice_viavel, color="#2ca02c", ls="--", lw=1, label="Limiar de viabilidade")
    ax.axhline(cfg.decisao.limiar_indice_ressalvas, color="#d62728", ls=":", lw=1, label="Limiar com ressalvas")
    for i, v in enumerate(vals):
        ax.text(i, v + 0.01, fmt_num(v, 3), ha="center")
    ax.set_ylim(0, 1.15)
    ax.legend(loc="upper center", ncol=2, frameon=False)
    salvar(fig, "FIG-INDICE")

    if "monte_carlo" in resultado:
        mc = resultado["monte_carlo"]
        v = np.sort(mc["vpl"])
        fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 3.8))
        a1.hist(v, bins=60, color="#4c9be8", alpha=0.85)
        a1.axvline(0, color="#d62728", lw=1)
        a1.axvline(np.mean(v), color="black", ls="--", lw=1, label="Média")
        a1.xaxis.set_major_formatter(reais)
        a1.set_xlabel("VPL (R$)")
        a1.set_ylabel("Frequência")
        a1.legend()
        a2.plot(v, np.arange(1, len(v) + 1) / len(v), color="#333333")
        a2.axvline(0, color="#d62728", lw=1)
        a2.xaxis.set_major_formatter(reais)
        a2.set_xlabel("VPL (R$)")
        a2.set_ylabel("Probabilidade acumulada")
        salvar(fig, "FIG-VPL")

        cv = mc["convergencia"]
        fig, ax = plt.subplots(figsize=(8, 4))
        ax.plot(cv["Iterações"], cv["Média do VPL"], color="#1f77b4", label="Média acumulada")
        ax.fill_between(cv["Iterações"], cv["IC inferior"], cv["IC superior"], color="#1f77b4", alpha=0.2, label="IC 95%")
        ax.set_xscale("log")
        ax.yaxis.set_major_formatter(reais)
        ax.set_xlabel("Iterações (escala logarítmica)")
        ax.set_ylabel("VPL médio (R$)")
        ax.legend()
        salvar(fig, "FIG-CONVERGENCIA")

    if "sensibilidade" in resultado:
        s = resultado["sensibilidade"].iloc[::-1]
        fig, ax = plt.subplots(figsize=(8, 4.2))
        y = np.arange(len(s))
        ax.barh(y, s["ΔVPL (-10%)"], color="#f2a65a", label="−10%")
        ax.barh(y, s["ΔVPL (+10%)"], color="#4c9be8", label="+10%")
        ax.set_yticks(y)
        ax.set_yticklabels(s["Rótulo"])
        ax.axvline(0, color="black", lw=0.8)
        ax.xaxis.set_major_formatter(reais)
        ax.set_xlabel("Variação do VPL (R$)")
        ax.legend()
        salvar(fig, "FIG-TORNADO")

    if alternativas is not None and len(alternativas):
        a = alternativas.sort_values("TOPSIS")
        fig, ax = plt.subplots(figsize=(8, 3.5))
        ax.barh(a["Alternativa"], a["TOPSIS"], color="#2ca02c")
        for i, v in enumerate(a["TOPSIS"]):
            ax.text(v + 0.01, i, fmt_num(v, 3), va="center")
        ax.set_xlim(0, 1.1)
        ax.set_xlabel("Proximidade relativa à solução ideal")
        salvar(fig, "FIG-TOPSIS")

    lg = resultado["lean_green"]
    fig, ax = plt.subplots(figsize=(8, 3.5))
    itens = {"Ração excedente": _f(lg["racao_excedente_rs"]), "Mortalidade evitável": _f(lg["mortalidade_evitavel_rs"]),
             "Energia excedente": _f(lg["energia_excedente_rs"])}
    ax.barh(list(itens), list(itens.values()), color="#f2a65a")
    for i, v in enumerate(itens.values()):
        ax.text(v, i, " " + fmt_rs(v), va="center")
    ax.set_xlabel("R$/ano")
    salvar(fig, "FIG-LEANGREEN")
    return saida


# ---------------------------------------------------------------------------------------------
# Exportações
# ---------------------------------------------------------------------------------------------
def resumo_executivo(resultado: Dict[str, Any], cfg: EVTEASConfig) -> str:
    m = marcadores(resultado, cfg).set_index("Marcador")["Valor"]
    g = lambda k: m.get(k, "n/d")
    linhas = [
        f"EVTEAS-Py {VERSAO} — {cfg.projeto}",
        f"Classificação: {g('R-I08')} | Índice EVTEAS: {g('R-I06')} | Vetos: {g('R-I09')} | Ressalvas: {g('R-I10')}",
        f"Técnica: produção {g('R-T05')} kg/ano; OEE {g('R-T08')}; FCR {fmt_num(cfg.tecnico.fcr, 2)}",
        f"Econômica: investimento {g('R-E02')}; VPL {g('R-E04')}; TIR {g('R-E06')}; payback descontado {g('R-E08')} meses",
        f"Beneficiário: VPL {g('R-E05')}; TIR {g('R-E07')}",
        f"Ambiental: PH azul {g('R-A03')} m³/t; PH cinza {g('R-A04')} m³/t; {g('R-A13')} kgCO2e/kg; conformidade {g('R-A15')}/10",
        f"Social: RVL {g('R-S01')}; LSO {g('R-S05')}; renda {g('R-S03')} ({g('R-S04')} SM)",
    ]
    if "monte_carlo" in resultado:
        linhas.append(f"Monte Carlo: VPL médio {g('R-M02')}; P(VPL>0) {g('R-M06')}")
    linhas.append(f"Verificação: {g('R-W01')} invariantes aprovados")
    return "\n".join(linhas)


def _jsonavel(x):
    if isinstance(x, dict):
        return {str(k): _jsonavel(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_jsonavel(v) for v in x]
    if isinstance(x, pd.DataFrame):
        return x.to_dict(orient="records")
    if isinstance(x, np.ndarray):
        return x.tolist() if x.size <= 500 else {"tamanho": int(x.size), "media": float(np.nanmean(x))}
    if isinstance(x, (np.floating, float)):
        return float(x) if math.isfinite(float(x)) else None
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, np.bool_):
        return bool(x)
    return x


def exportar_excel(resultado: Dict[str, Any], cfg: EVTEASConfig, caminho: str,
                   alternativas: Optional[pd.DataFrame] = None, testes: Optional[pd.DataFrame] = None,
                   configs_alternativas: Optional[Dict[str, EVTEASConfig]] = None) -> str:
    """Workbook com premissas, resultados e registros da verificação."""
    abas: Dict[str, pd.DataFrame] = {}
    abas["Resumo"] = pd.DataFrame({"Resumo executivo": resumo_executivo(resultado, cfg).split("\n")})
    abas["Marcadores"] = marcadores(resultado, cfg, alternativas, testes, configs_alternativas)
    abas["Entradas"] = tabela_entradas(cfg)
    abas["Curva crescimento"] = resultado["curva_crescimento"]
    abas["DRE mensal"] = resultado["dre_mensal"]
    abas["DRE anual"] = resultado["dre_anual"]
    abas["Fluxo de caixa"] = pd.DataFrame({"Mês": np.arange(len(resultado["fluxo_caixa"])),
                                           "Fluxo": resultado["fluxo_caixa"], "Acumulado": np.cumsum(resultado["fluxo_caixa"])})
    abas["Sobras"] = pd.DataFrame(list(resultado["sobras"].items()), columns=["Destinação", "R$"])
    A = resultado["ambiental"]
    abas["Ambiental"] = pd.DataFrame([(k, _f(v)) for k, v in A.items() if not isinstance(v, (pd.DataFrame, list, str))
                                      and np.size(v) == 1 and not isinstance(np.atleast_1d(v)[0], str)], columns=["Indicador", "Valor"])
    abas["Conformidade"] = A["checklist"]
    S = resultado["social"]
    abas["Social"] = pd.DataFrame([(k, _f(v)) for k, v in S.items() if not isinstance(v, dict) and np.size(v) == 1],
                                  columns=["Indicador", "Valor"])
    abas["Radar social"] = pd.DataFrame([(k, _f(v)) for k, v in S["radar"].items()], columns=["Eixo", "Valor"])
    abas["LSO componentes"] = pd.DataFrame(list(S["componentes_lso"].items()), columns=["Componente", "Valor"])
    abas["Governança"] = resultado["governanca"]["checklist"]
    abas["ODS"] = resultado["ods"]
    abas["Lean-Green"] = pd.DataFrame([(k, _f(v)) for k, v in resultado["lean_green"].items()], columns=["Indicador", "Valor"])
    abas["KPIs normalizados"] = resultado["normalizacao"]
    if "monte_carlo" in resultado:
        mc = resultado["monte_carlo"]
        abas["MC estatísticas"] = pd.DataFrame([(k, v) for k, v in mc["estatisticas"].items()], columns=["Estatística", "Valor"])
        abas["MC amostra"] = mc["amostra"].head(10000)
        abas["MC convergência"] = mc["convergencia"]
        abas["Importância"] = mc["importancia"]
    if "sensibilidade" in resultado:
        abas["Sensibilidade"] = resultado["sensibilidade"]
        abas["Valores críticos"] = resultado["valores_criticos"]
        abas["Cenários"] = resultado["cenarios"]
    if alternativas is not None:
        abas["Alternativas TOPSIS"] = alternativas
    abas["Verificação"] = resultado["vv"]["tabela"]
    if testes is not None:
        abas["Testes"] = testes
    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        for nome, df in abas.items():
            df.to_excel(w, sheet_name=nome[:31], index=False)
    return caminho


def exportar_json(resultado: Dict[str, Any], caminho: str) -> str:
    dados = {k: v for k, v in resultado.items() if k not in ("dre_mensal",)}
    if "monte_carlo" in dados:
        dados["monte_carlo"] = {k: v for k, v in dados["monte_carlo"].items() if k not in ("amostra", "entradas", "vpl", "indice", "vpl_beneficiario")}
    Path(caminho).write_text(json.dumps(_jsonavel(dados), ensure_ascii=False, indent=2), encoding="utf-8")
    return caminho


def exportar_markdown(resultado: Dict[str, Any], cfg: EVTEASConfig, caminho: str,
                      alternativas: Optional[pd.DataFrame] = None, testes: Optional[pd.DataFrame] = None,
                      configs_alternativas: Optional[Dict[str, EVTEASConfig]] = None) -> str:
    partes = [f"# Resultados do EVTEAS-Py {VERSAO}", "", f"**Projeto:** {cfg.projeto}", "", "```",
              resumo_executivo(resultado, cfg), "```", "", "## Marcadores para o texto da dissertação", ""]
    mk = marcadores(resultado, cfg, alternativas, testes, configs_alternativas)
    partes.append("| Marcador | Descrição | Valor | Seção |")
    partes.append("|---|---|---|---|")
    for _, r in mk.iterrows():
        partes.append(f"| [{r['Marcador']}] | {r['Descrição']} | {str(r['Valor']).replace('|', '/')} | {r['Seção']} |")
    for nome, df in tabelas_dissertacao(resultado, cfg, alternativas).items():
        partes += ["", f"## {nome}", "", df.to_markdown(index=False) if hasattr(df, "to_markdown") and _tem_tabulate() else df.to_string(index=False)]
    Path(caminho).write_text("\n".join(partes), encoding="utf-8")
    return caminho


def _tem_tabulate() -> bool:
    try:
        import tabulate  # noqa: F401
        return True
    except ImportError:
        return False


MARCADOR_DOCX = re.compile(r"\[((?:R|E)-[A-Z0-9]+|FIG-[A-Z]+)\]")


def _substituir_no_paragrafo(p, padrao, funcao) -> None:
    """Substitui ocorrências do padrão no parágrafo mesmo quando o marcador está dividido entre
    vários trechos (runs) do Word; o valor herda a formatação do trecho onde o marcador começa."""
    texto = "".join(r.text for r in p.runs)
    for m in reversed(list(padrao.finditer(texto))):
        novo = funcao(m)
        if novo is None:
            continue
        ini, fim = m.span()
        pos, afetados = 0, []
        for r in p.runs:
            a, b = pos, pos + len(r.text)
            if b > ini and a < fim:
                afetados.append((r, a))
            pos = b
        if not afetados:
            continue
        r0, a0 = afetados[0]
        if len(afetados) == 1:
            r0.text = r0.text[:ini - a0] + novo + r0.text[fim - a0:]
        else:
            r0.text = r0.text[:ini - a0] + novo
            for r, _ in afetados[1:-1]:
                r.text = ""
            rl, al = afetados[-1]
            rl.text = rl.text[fim - al:]
        r0.font.highlight_color = None


def preencher_docx(entrada: str, saida: str, marcadores_df: pd.DataFrame, figuras: Optional[Dict[str, str]] = None,
                   largura_cm: float = 15.0) -> Dict[str, Any]:
    """Substitui os marcadores [E-xx] e [R-xx] de um .docx pelos valores da simulação (mantendo a
    formatação do trecho e removendo o realce) e insere as figuras nos parágrafos com [FIG-...].
    Marcadores sem valor nesta execução permanecem no texto e são listados. Requer python-docx."""
    from docx import Document
    from docx.shared import Cm
    valores = dict(zip(marcadores_df["Marcador"], marcadores_df["Valor"].astype(str)))
    doc = Document(entrada)
    cont = {"trocados": 0, "imagens": 0}
    ausentes = set()

    def paragrafos(container):
        for p in container.paragraphs:
            yield p
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from paragrafos(cell)

    def valor(m):
        chave = m.group(1)
        if chave.startswith("FIG-"):
            return None
        if chave in valores:
            cont["trocados"] += 1
            return valores[chave]
        ausentes.add(chave)
        return None

    for p in paragrafos(doc):
        if "[" not in p.text:
            continue
        _substituir_no_paragrafo(p, MARCADOR_DOCX, valor)
        figs = re.findall(r"\[(FIG-[A-Z]+)\]", p.text)
        for chave in figs:
            if figuras and chave in figuras and os.path.exists(figuras[chave]):
                _substituir_no_paragrafo(p, re.compile(re.escape(f"[{chave}]")), lambda m: "")
                p.add_run().add_picture(figuras[chave], width=Cm(largura_cm))
                cont["imagens"] += 1
            else:
                ausentes.add(chave)
    doc.save(saida)
    return {"marcadores_substituidos": cont["trocados"], "marcadores_sem_valor": sorted(ausentes),
            "figuras_inseridas": cont["imagens"], "arquivo": saida}


def exportar_tudo(resultado: Dict[str, Any], cfg: EVTEASConfig, pasta: str, figuras: bool = True,
                  alternativas: Optional[pd.DataFrame] = None, testes: Optional[pd.DataFrame] = None,
                  configs_alternativas: Optional[Dict[str, EVTEASConfig]] = None) -> Dict[str, str]:
    """Gera todas as saídas na pasta e um .zip com tudo."""
    from .config import salvar_config
    Path(pasta).mkdir(parents=True, exist_ok=True)
    arq = {
        "excel": exportar_excel(resultado, cfg, os.path.join(pasta, "EVTEAS_resultados.xlsx"), alternativas, testes,
                                configs_alternativas),
        "json": exportar_json(resultado, os.path.join(pasta, "EVTEAS_resultados.json")),
        "markdown": exportar_markdown(resultado, cfg, os.path.join(pasta, "EVTEAS_relatorio.md"), alternativas, testes,
                                      configs_alternativas),
        "config": str(salvar_config(cfg, os.path.join(pasta, "EVTEAS_premissas.json"))),
    }
    mk = marcadores(resultado, cfg, alternativas, testes, configs_alternativas)
    arq["marcadores"] = os.path.join(pasta, "EVTEAS_marcadores.csv")
    mk.to_csv(arq["marcadores"], index=False, encoding="utf-8-sig", sep=";")
    Path(os.path.join(pasta, "EVTEAS_resumo.txt")).write_text(resumo_executivo(resultado, cfg), encoding="utf-8")
    arq["resumo"] = os.path.join(pasta, "EVTEAS_resumo.txt")
    if figuras:
        for k, v in gerar_figuras(resultado, cfg, os.path.join(pasta, "figuras"), alternativas).items():
            arq[k] = v
    zipp = os.path.join(pasta, "EVTEAS_saidas.zip")
    with zipfile.ZipFile(zipp, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in arq.items():
            z.write(v, arcname=os.path.relpath(v, pasta))
    arq["zip"] = zipp
    return arq


def comparar_pesos(r_preset: Dict[str, Any], r_likert: Dict[str, Any]) -> pd.DataFrame:
    """Marcadores [R-Lxx]: efeito dos pesos Likert dos especialistas (obtidos nos questionários,
    após a simulação) sobre o índice e a classificação, em comparação com o preset."""
    ip, il = r_preset["indice"], r_likert["indice"]
    pl = il["pesos"]
    w = pl["dimensoes"]
    linhas = [
        ("R-L01", "Pesos dimensionais Likert (T/E/A/S)",
         f"{fmt_num(w['tecnica'], 3)} / {fmt_num(w['economica'], 3)} / {fmt_num(w['ambiental'], 3)} / {fmt_num(w['social'], 3)}"),
        ("R-L02", "Concordância média entre os especialistas", fmt_num(pl.get("concordancia_media"), 3)),
        ("R-L03", "Índice EVTEAS: preset → Likert",
         f"{fmt_num(_f(ip['indice_evteas']), 3)} → {fmt_num(_f(il['indice_evteas']), 3)}"),
        ("R-L04", "Classificação: preset → Likert",
         f"{r_preset['decisao']['classificacao']} → {r_likert['decisao']['classificacao']}"),
    ]
    conc = pl.get("concordancia_por_kpi", {})
    if conc:
        baixos = [KPIS[k][1] for k, v in sorted(conc.items(), key=lambda kv: kv[1])[:3]]
        linhas.append(("R-L05", "KPIs de menor concordância entre os especialistas", "; ".join(baixos)))
    return pd.DataFrame([(m, d, v, "4.7") for m, d, v in linhas], columns=["Marcador", "Descrição", "Valor", "Seção"])
